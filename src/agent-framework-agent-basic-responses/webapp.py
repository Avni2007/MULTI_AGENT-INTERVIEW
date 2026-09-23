# Copyright (c) Microsoft. All rights reserved.
# Employability AI — FastAPI web application
# Runs independently of ResponsesHostServer (main.py is untouched).

import io
import json
import logging
import os
import re
import textwrap
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse

# ── Azure Document Intelligence ──────────────────────────────────────────────
from azure.ai.formrecognizer import DocumentAnalysisClient
from azure.core.credentials import AzureKeyCredential
from azure.core.exceptions import HttpResponseError

# ── Azure Identity ───────────────────────────────────────────────────────────
from azure.identity import DefaultAzureCredential

# ── Agent Framework ──────────────────────────────────────────────────────────
from agent_framework import Agent
from agent_framework.foundry import FoundryChatClient

# ── Async support ────────────────────────────────────────────────────────────
import asyncio

# ─────────────────────────────────────────────────────────────────────────────
# Logging
# ─────────────────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
)
logger = logging.getLogger("employability-ai")

# ─────────────────────────────────────────────────────────────────────────────
# Environment
# ─────────────────────────────────────────────────────────────────────────────
_local_env = Path(__file__).parent / ".env"
if _local_env.exists():
    load_dotenv(_local_env)
load_dotenv()

_REQUIRED_VARS = [
    "FOUNDRY_PROJECT_ENDPOINT",
    "AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT",
    "AZURE_DOCUMENT_INTELLIGENCE_KEY",
]
_missing = [v for v in _REQUIRED_VARS if not os.environ.get(v)]
if _missing:
    raise RuntimeError(
        f"Missing required environment variables: {', '.join(_missing)}. "
        "Check your .env file."
    )

MODEL_NAME = os.getenv("AZURE_AI_MODEL_DEPLOYMENT_NAME") or os.getenv("FOUNDRY_MODEL_NAME")
if not MODEL_NAME:
    raise RuntimeError(
        "Model deployment name is not configured. "
        "Set AZURE_AI_MODEL_DEPLOYMENT_NAME or FOUNDRY_MODEL_NAME in .env."
    )

# ─────────────────────────────────────────────────────────────────────────────
# Azure clients (initialised once at startup)
# ─────────────────────────────────────────────────────────────────────────────
document_client = DocumentAnalysisClient(
    endpoint=os.environ["AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT"],
    credential=AzureKeyCredential(os.environ["AZURE_DOCUMENT_INTELLIGENCE_KEY"]),
)

foundry_client = FoundryChatClient(
    project_endpoint=os.environ["FOUNDRY_PROJECT_ENDPOINT"],
    model=MODEL_NAME,
    credential=DefaultAzureCredential(),
)

# ─────────────────────────────────────────────────────────────────────────────
# Agent
# ─────────────────────────────────────────────────────────────────────────────
SYSTEM_INSTRUCTIONS = textwrap.dedent("""\
    You are a professional Resume Analysis Agent for the "Employability AI" platform.

    Analyze the candidate's resume using ONLY evidence found in the provided resume text.

    SCORING RUBRIC (total = 100):
      Structure       /10  — Sections, formatting, logical flow
      Skills          /15  — Relevance, depth, variety of technical/soft skills
      Experience      /20  — Roles, responsibilities, impact, tenure
      Projects        /15  — Complexity, outcomes, technologies used
      Achievements    /10  — Quantifiable wins, awards, recognition
      Education       /10  — Degrees, institutions, relevance
      Job Alignment   /15  — Fit against job description (0 if none provided)
      Clarity & Lang  /5   — Writing quality, grammar, conciseness

    AGENT RULES:
    - Score ONLY evidence found in the resume.
    - NEVER invent skills, experience, employers, achievements, qualifications, or metrics.
    - If information is absent, use "Not found" as the reason.
    - Clearly separate resume facts from recommendations.
    - If a job description is provided, identify requirements met and gaps.
    - Suggested improved bullets must NOT invent metrics; only restructure existing ones.
    - Return ONLY valid JSON — no markdown fences, no prose, no extra keys.

    OUTPUT FORMAT (strict JSON):
    {
      "overall_score": <integer 0-100>,
      "category_scores": {
        "structure":        {"score": <0-10>,  "reason": "<string>"},
        "skills":           {"score": <0-15>,  "reason": "<string>"},
        "experience":       {"score": <0-20>,  "reason": "<string>"},
        "projects":         {"score": <0-15>,  "reason": "<string>"},
        "achievements":     {"score": <0-10>,  "reason": "<string>"},
        "education":        {"score": <0-10>,  "reason": "<string>"},
        "job_alignment":    {"score": <0-15>,  "reason": "<string>"},
        "clarity_language": {"score": <0-5>,   "reason": "<string>"}
      },
      "strengths":          ["<string>", ...],
      "improvements":       ["<string>", ...],
      "missing_information":["<string>", ...],
      "weak_bullets":       [{"original": "<string>", "improved": "<string>"}, ...],
      "job_matches":        ["<string>", ...],
      "job_gaps":           ["<string>", ...]
    }
""")

agent = Agent(
    client=foundry_client,
    instructions=SYSTEM_INSTRUCTIONS,
    default_options={"store": False},
)

# Thread-pool for blocking Document Intelligence SDK calls (sync SDK)
_doc_executor = ThreadPoolExecutor(max_workers=4)

# ─────────────────────────────────────────────────────────────────────────────
# Constants
# ─────────────────────────────────────────────────────────────────────────────
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024   # 10 MB
MAX_JD_CHARS = 5_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc"}

# ─────────────────────────────────────────────────────────────────────────────
# FastAPI app
# ─────────────────────────────────────────────────────────────────────────────
app = FastAPI(title="Employability AI", version="1.0.0")

HERE = Path(__file__).parent


# ── Helper: extract text via Document Intelligence (blocking SDK) ─────────────

def _extract_docx_fallback(file_bytes: bytes) -> str:
    """Fallback DOCX extraction using python-docx (no Azure call)."""
    try:
        from docx import Document as DocxDocument  # type: ignore
        doc = DocxDocument(io.BytesIO(file_bytes))
        lines = []
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                lines.append(t)
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t and t not in lines:
                        lines.append(t)
        return "\n".join(lines).strip()
    except Exception as exc:
        logger.warning("DOCX fallback extraction failed: %s", exc)
        return ""


def _extract_text_blocking(file_bytes: bytes, filename: str) -> str:
    """Run Document Intelligence synchronously — called from thread pool.
    Falls back to python-docx for DOCX files if Document Intelligence returns
    no text (e.g. simple files without proper OCR content markers).
    """
    suffix = Path(filename).suffix.lower()
    logger.info(
        "Sending %d bytes (%s) to Azure Document Intelligence …",
        len(file_bytes), suffix,
    )
    try:
        poller = document_client.begin_analyze_document(
            model_id="prebuilt-read",
            document=io.BytesIO(file_bytes),
        )
        result = poller.result()
    except HttpResponseError as exc:
        logger.error("Document Intelligence HTTP error: %s", exc)
        raise RuntimeError(
            f"Azure Document Intelligence error: {exc.message or str(exc)}"
        )

    lines = []
    for page in result.pages:
        for line in page.lines:
            lines.append(line.content)
    text = "\n".join(lines).strip()
    logger.info("Azure Document Intelligence extracted %d characters.", len(text))

    # Fallback for DOCX: if Document Intelligence extracted nothing, try python-docx
    if not text and suffix in (".docx", ".doc"):
        logger.info("Document Intelligence returned 0 chars for DOCX — using python-docx fallback.")
        text = _extract_docx_fallback(file_bytes)
        logger.info("python-docx fallback extracted %d characters.", len(text))

    return text


async def _extract_text(file_bytes: bytes, filename: str) -> str:
    """Async wrapper — offloads blocking SDK call to thread pool."""
    loop = asyncio.get_event_loop()
    try:
        return await loop.run_in_executor(
            _doc_executor,
            _extract_text_blocking,
            file_bytes,
            filename,
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=502, detail=str(exc))


# ── Helper: call agent (async-native) ────────────────────────────────────────

async def _call_agent(resume_text: str, job_description: str | None) -> dict:
    """Build prompt and call Foundry GPT-4.1-mini via Agent Framework."""
    prompt_parts = ["RESUME TEXT:\n" + resume_text]
    if job_description and job_description.strip():
        prompt_parts.append(
            "\nJOB DESCRIPTION:\n" + job_description.strip()[:MAX_JD_CHARS]
        )
    else:
        prompt_parts.append(
            "\nNo job description provided — set job_alignment score to 0 with reason 'No job description provided'."
        )

    prompt = "\n\n".join(prompt_parts)
    logger.info("Sending prompt (%d chars) to Foundry …", len(prompt))

    try:
        # agent.run() returns an Awaitable[AgentResponse]
        response = await agent.run(prompt)
    except Exception as exc:
        logger.error("Foundry agent error: %s", exc, exc_info=True)
        raise HTTPException(
            status_code=502,
            detail=f"Foundry model error: {exc}",
        )

    # Extract text from AgentResponse
    raw: str = ""
    if hasattr(response, "text"):
        raw = response.text or ""
    elif hasattr(response, "content"):
        raw = str(response.content)
    elif hasattr(response, "output_text"):
        raw = str(response.output_text)
    else:
        raw = str(response)

    logger.info("Raw model response (first 600 chars): %s", raw[:600])

    # Robust JSON parsing
    raw = raw.strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        pass

    # Strip markdown code fences if present
    cleaned = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned.rstrip())
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        pass

    # Extract outermost JSON object { ... }
    match = re.search(r"(\{.*\})", raw, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(1))
        except json.JSONDecodeError:
            pass

    logger.error("JSON parse error on raw:\n%s", raw)
    raise HTTPException(
        status_code=502,
        detail="Model returned invalid JSON. Ensure the model is instructed to return only JSON.",
    )


# ─────────────────────────────────────────────────────────────────────────────
# Routes
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def serve_ui():
    """Serve the frontend HTML."""
    html_path = HERE / "index.html"
    if not html_path.exists():
        raise HTTPException(status_code=404, detail="index.html not found.")
    return HTMLResponse(content=html_path.read_text(encoding="utf-8"))


@app.post("/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    job_description: str = Form(default=""),
):
    """
    Accept a PDF or DOCX resume, extract text via Azure Document Intelligence,
    analyse with Foundry GPT-4.1-mini, return structured JSON.
    """
    # ── Validate file type ────────────────────────────────────────────────────
    filename = file.filename or "upload"
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unsupported file type '{suffix}'. "
                "Please upload a PDF or DOCX file."
            ),
        )

    # ── Read file bytes ───────────────────────────────────────────────────────
    file_bytes = await file.read()
    if len(file_bytes) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=(
                f"File too large ({len(file_bytes) // 1024} KB). "
                "Maximum allowed size is 10 MB."
            ),
        )

    logger.info(
        "Received '%s' (%d bytes), job_description length: %d",
        filename, len(file_bytes), len(job_description),
    )

    # ── Extract text ──────────────────────────────────────────────────────────
    resume_text = await _extract_text(file_bytes, filename)
    if not resume_text:
        raise HTTPException(
            status_code=422,
            detail=(
                "Could not extract any text from the uploaded file. "
                "Ensure the document contains selectable text (not a scanned image)."
            ),
        )

    # ── Analyse with Foundry agent ────────────────────────────────────────────
    result = await _call_agent(resume_text, job_description or None)

    return JSONResponse(content=result)


@app.get("/health")
async def health():
    """Quick health/readiness check."""
    return {"status": "ok", "model": MODEL_NAME}


# ─────────────────────────────────────────────────────────────────────────────
# Entry point
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn

    port = int(os.getenv("PORT", "8080"))
    logger.info("Starting Employability AI on http://localhost:%d", port)
    uvicorn.run(app, host="0.0.0.0", port=port)
